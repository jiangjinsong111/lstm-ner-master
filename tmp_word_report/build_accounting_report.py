from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from pathlib import Path
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm

OUT_DIR = Path('tmp_word_report/output')
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT = OUT_DIR / '合肥工业大学会计学专业培养方案对标调研分析报告_语言修订版.docx'
CHART = OUT_DIR / '图1_主课程体系学分位置.png'

# ---------- chart ----------
labels = ['合肥工业大学（现行）','北京航空航天大学','中南大学','华中科技大学','浙江大学','北京理工大学','哈尔滨工业大学','厦门大学']
values = [165, 163, 159.5, 157.5, 155, 152, 150, 140]
font_candidates = [
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
    '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc',
    '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc'
]
font_path = next((p for p in font_candidates if Path(p).exists()), None)
font_prop = fm.FontProperties(fname=font_path) if font_path else None
plt.figure(figsize=(10.8, 5.5))
ax = plt.gca()
y = list(range(len(labels)))
bars = ax.barh(y, values, color=['#3a3a3a'] + ['#c9c9c9'] * 7, edgecolor='#666666', linewidth=0.7)
ax.invert_yaxis()
ax.set_xlim(136, 168)
ax.set_yticks(y)
ax.set_yticklabels(labels, fontproperties=font_prop, fontsize=10)
ax.set_xlabel('纳入比较的主课程体系学分（中国高校口径）', fontproperties=font_prop, fontsize=10)
ax.axvspan(151.0, 158.5, color='#eeeeee', zorder=0)
ax.axvline(153.9, color='#555555', linestyle='--', linewidth=1)
for i, (bar, val) in enumerate(zip(bars, values)):
    ax.text(val + 0.35, bar.get_y() + bar.get_height()/2, f'{val:g}', va='center', fontsize=9)
ax.text(159.3, 6.3, '7校四分位区间 151.0—158.5', fontproperties=font_prop, fontsize=9)
ax.text(159.3, 6.8, '7校均值 153.9', fontproperties=font_prop, fontsize=9)
ax.spines['top'].set_visible(False)
ax.spines['right'].set_visible(False)
plt.tight_layout()
plt.savefig(CHART, dpi=220, bbox_inches='tight')
plt.close()

# ---------- helpers ----------
def set_east_asia_font(run, name='宋体', size=12, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['val','sz','space','color']:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def add_body(doc, text, first_indent=True, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    if first_indent:
        pf.first_line_indent = Cm(0.85)
    r = p.add_run(text)
    set_east_asia_font(r, '宋体', 12)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_east_asia_font(r, '黑体', 14, False)
    return p


def add_subpara(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(0.85)
    r1 = p.add_run(label)
    set_east_asia_font(r1, '宋体', 12, False)
    r2 = p.add_run(text)
    set_east_asia_font(r2, '宋体', 12, False)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_east_asia_font(r, '宋体', 10.5, True)
    return p


def add_note(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_east_asia_font(r, '宋体', 8.5)
    return p

# ---------- document ----------
doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.15)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.45)
sec.right_margin = Cm(2.25)
sec.header_distance = Cm(1.0)
sec.footer_distance = Cm(1.0)

styles = doc.styles
normal = styles['Normal']
normal.font.name = '宋体'
normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5

# footer page number
footer_p = sec.footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer_p.add_run()
set_east_asia_font(run, '宋体', 10)
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

# title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run('合肥工业大学会计学专业培养方案对标调研分析报告')
set_east_asia_font(r, '黑体', 16, False)

add_h1(doc, '一、调研范围、比较口径与数据处理')
add_body(doc, '本次调研围绕培养目标、毕业要求、课程体系、实践教学、人才培养模式和教育教学改革举措六个方面展开。国内选取北京理工大学、北京航空航天大学、华中科技大学、厦门大学、哈尔滨工业大学、浙江大学和中南大学，共涉及7所高校、10份会计学相关培养方案，方案版本集中在2019—2023年。厦门大学包含会计学、注册会计师、国际会计和CIMA四个方向，因此学校数量按7所统计，文档数量按10份统计。国外分别选取亚洲、欧洲和美洲各1所高校，即新加坡南洋理工大学、英国伦敦政治经济学院和美国伊利诺伊大学厄巴纳—香槟分校。相关资料主要来自学校官网项目页面和2026—2027课程目录，检索截至2026年8月2日。')
add_body(doc, '国内高校的学分统计采用主课程体系比较值。培养方案正文或课程结构主表明确列入的课内课程和实践教学学分计入比较值；另行列出的第二课堂、课外活动和附加模块，如未计入主表合计，则单独说明；已包含在主表合计中的项目不再重复计算。表1列出各校的取值依据。统计以学校为单位，厦门大学四个方向在7校统计中计1次。国外高校采用的LSE unit、semester hour和AU分别属于不同学分制度，本报告仅呈现各项目内部的课程结构，不作直接换算。')

add_caption(doc, '表1  国内7校会计学相关方案主课程体系学分比较值及认定依据')
table1_data = [
['学校','版本','比较值','比较值认定依据（学分）'],
['北京航空航天大学','2022','163','数字化与智能会计方向正文规定毕业总学分163；比较值取163'],
['中南大学','2023','159.5','课程结构主表合计159.5，已含课外研学4；比较值取159.5'],
['华中科技大学','2021','157.5','课内课程与集中实践合计157.5；另要求课外5，后者不计入'],
['浙江大学','2019','155','文件写作155+5.5+6+8；155为主课程体系，另列19.5不计入比较值'],
['北京理工大学','2019','152','正文规定最低152，课程计划表合计149；以正文为准，比较值取152'],
['哈尔滨工业大学','2022','150','正文规定修满150并完成论文答辩方可毕业；比较值取150'],
['厦门大学','2019','140','会计学、注册会计师、国际会计、CIMA四方向均为140；学校统计只计1次，比较值取140']]
t = doc.add_table(rows=len(table1_data), cols=4)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [Cm(3.4), Cm(1.7), Cm(1.8), Cm(10.0)]
for i,row in enumerate(t.rows):
    for j,cell in enumerate(row.cells):
        cell.width = widths[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 3 else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        rr = p.add_run(table1_data[i][j])
        set_east_asia_font(rr, '宋体', 8.5, i==0)
        set_cell_border(cell, top={'val':'single','sz':'6','color':'000000'}, bottom={'val':'single','sz':'6','color':'000000'}, left={'val':'single','sz':'4','color':'000000'}, right={'val':'single','sz':'4','color':'000000'})
set_repeat_table_header(t.rows[0])
add_note(doc, '注：比较值用于本次7校样本均值、中位数和四分位数的计算，不涵盖学生全部毕业条件。浙江大学另列的19.5学分由“+5.5学分”、跨专业与国际化模块6学分、第二至第四课堂8学分构成；华中科技大学另列课外5学分。相关数据均根据各校培养方案正文和课程结构表整理。')

add_body(doc, '按照上述口径，7校主课程体系比较值合计1077学分，平均值为153.9学分，中位数为155学分。比较值最低为140学分，最高为163学分，极差为23学分。采用线性插值计算，第一四分位数为151.0学分，第三四分位数为158.5学分，中间50%的学校分布在151.0—158.5学分之间。按本次7校样本总体计算，标准差为7.0学分。如按10份文档统计，厦门大学四个方向会形成重复权重，平均值为149.7学分，中位数为151学分。因此，后文统一采用7校统计口径。')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(str(CHART), width=Cm(15.6))
add_caption(doc, '图1  国内7校及合肥工业大学主课程体系学分位置')
add_note(doc, '注：合肥工业大学现行方案共165学分，其中理论课程134.5学分、实践教学30.5学分，与7校主课程体系比较值采用同类统计范围。按该口径计算，合肥工业大学比7校平均值153.9学分多11.1学分，相差7.2%；比样本最高值163学分多2学分。')

add_h1(doc, '二、国内高校六维调研结果')
add_subpara(doc, '（一）培养目标。', '国内样本的培养目标主要涉及会计专业知识、数据分析和决策支持。北京航空航天大学将信息技术、大数据分析和智能技术纳入培养目标，提出运用财务、业务和社会经济数据支持管理决策；华中科技大学大数据与商业分析方向同时列出会计审计知识、程序设计和数据分析能力。两所学校均结合业务问题和决策任务描述相关能力。')
add_subpara(doc, '（二）毕业要求。', '国内样本中的毕业要求通常涵盖专业知识、问题分析、实践应用、沟通协作、职业规范、创新和持续学习等内容。华中科技大学列出程序设计、管理信息系统开发和大数据分析方法；北京航空航天大学将运用计算机和大数据处理技术解决会计实务与决策问题纳入毕业要求。部分方案还将课程项目、实习和毕业成果作为相关能力的评价载体。')
add_subpara(doc, '（三）课程体系。', '财务会计、管理会计、成本会计、审计、财务管理、税法和会计信息系统是国内样本中较为稳定的专业核心课程。以7所学校为统计单位，7所学校均开设计算机、信息系统或数据分析基础课程，占100%；北京理工大学、北京航空航天大学、华中科技大学、浙江大学和中南大学明确开设Python、C或C++课程，共5所，占71.4%。在程序设计课程之后，部分学校继续设置数据库、业务分析或智能会计课程。')
add_subpara(doc, '（四）实践教学。', '不同学校采用不同方式组织和计量实践教学。北京航空航天大学的社会课堂（生产实习）为5学分、320学时；华中科技大学集中性实践教学为17学分、34周；哈尔滨工业大学的生产实习和毕业实习各3学分、各3周。相关方案分别以学分、学时和周数说明实践要求，对实践任务、企业参与和成果形式的列示方式也有所不同。')
add_subpara(doc, '（五）人才培养模式。', '厦门大学在共同基础课程之外设置会计学、注册会计师、国际会计和CIMA四个方向，四个方向均为140学分；北京航空航天大学设置数字化与智能会计方向；浙江大学的课程结构包括宽口径基础、专业选修、跨专业模块和国际化模块。各校在共同核心课程、专业方向和跨学科选修空间方面采用了不同的组织方式。')
add_subpara(doc, '（六）教育教学改革举措。', '从国内样本看，相关改革主要涉及数智课程建设、案例与项目教学、校企合作和课程评价调整。部分学校按照程序设计、数据库或信息系统、大数据分析、智能财务或智能决策的顺序安排课程；部分学校通过专题课程或方向选修课程呈现相关内容。不同方案对课程先修关系和综合项目的安排各有侧重。')

add_h1(doc, '三、国外三所高校调研结果')
add_body(doc, '国外样本分别选取亚洲、欧洲和美洲各1所高校。南洋理工大学的项目资料主要呈现职业资格与长周期实习的衔接，伦敦政治经济学院的课程结构包括社会科学和数量方法，伊利诺伊大学厄巴纳—香槟分校将会计课程与数据科学课程组合设置。表2沿用各校官方学分口径，相关数字仅用于说明项目内部结构。')
add_caption(doc, '表2  亚洲、欧洲、美洲三所高校官方项目数据与培养机制')
table2_data = [
['区域','院校与项目','官方数据','培养机制'],
['亚洲','南洋理工大学\nAccountancy for Future Leaders','学制4年，授予荣誉学士；总计140 AU；30周Work-Study计15 AU','会计核心与可持续管理、分析能力一体化；实习机构须获ACRA认可，并衔接SCAQ部分考试减免'],
['欧洲','伦敦政治经济学院\nBSc Accounting and Finance','3年修读12 LSE units，另修LSE100；项目获ACCA、CIMA、ICAEW认可用于部分考试减免','会计与金融约占一半，配套经济学、数学、统计、计量经济学及校外选修；官方项目页未列必修长周期实习'],
['美洲','伊利诺伊大学厄巴纳—香槟分校\nAccountancy + Data Science','毕业最低124 semester hours；商科核心42、会计要求21、数据科学核心29—30；真实客户综合课程3学分','会计系与统计、计算机、信息学院和数学系协同；数据科学伦理、算法、建模和数据治理贯穿课程']]
t2 = doc.add_table(rows=len(table2_data), cols=4)
t2.alignment = WD_TABLE_ALIGNMENT.CENTER; t2.autofit = False
widths2 = [Cm(1.5), Cm(4.1), Cm(5.2), Cm(6.1)]
for i,row in enumerate(t2.rows):
    for j,cell in enumerate(row.cells):
        cell.width = widths2[j]; cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        rr = p.add_run(table2_data[i][j]); set_east_asia_font(rr,'宋体',8.2,i==0)
        set_cell_border(cell, top={'val':'single','sz':'6','color':'000000'}, bottom={'val':'single','sz':'6','color':'000000'}, left={'val':'single','sz':'4','color':'000000'}, right={'val':'single','sz':'4','color':'000000'})
set_repeat_table_header(t2.rows[0])
add_note(doc, '注：AU、LSE unit和semester hour分别是南洋理工大学、伦敦政治经济学院和美国高校使用的学分单位，三者不作直接换算。ACRA指新加坡会计与企业监管局，SCAQ指新加坡特许会计师资格考试。伊利诺伊大学124学时还包含通识教育等要求，表列模块不与中国高校学分直接比较。')
add_subpara(doc, '（一）亚洲：南洋理工大学。', '该项目将可持续管理、分析课程、专业会计和职业实践纳入同一培养路径。30周Work-Study实习计15 AU，实习单位须经新加坡会计与企业监管局认可，相关经历可计入新加坡特许会计师资格要求的实践年限。项目资料同时列明实习机构条件、学分认定和职业资格衔接方式。')
add_subpara(doc, '（二）欧洲：伦敦政治经济学院。', '项目学制为3年，学生修读12个LSE units，并完成LSE100。课程除会计和金融外，还包括经济学、数学、统计和计量经济学，学生可按规定选修校外课程。ACCA、CIMA和ICAEW的部分考试减免与学生的具体选课组合相衔接。')
add_subpara(doc, '（三）美洲：伊利诺伊大学厄巴纳—香槟分校。', '该项目毕业最低要求为124学时，其中会计课程21学时、数据科学核心课程29—30学时。数据科学部分包括数学基础、算法与数据结构、数据伦理和数据治理。BUS 301为3学分真实客户综合课程，课程内容包括商业问题处理、数据分析和客户展示。项目分别规定了会计课程、数据科学课程和综合课程的学分要求。')
add_body(doc, '按本报告选取的3所学校统计，3校均采用跨学科或宽口径课程结构，占100%；南洋理工大学和伊利诺伊大学将数据分析或数据科学纳入课程体系，占66.7%；南洋理工大学和伦敦政治经济学院明确提供职业资格考试减免，占66.7%；南洋理工大学的30周实习和伊利诺伊大学的3学分真实客户课程均计入培养要求，结构化实践占66.7%。其中，南洋理工大学设置长周期必修实习，占33.3%。上述比例仅反映本次3个案例的基本情况。')

add_h1(doc, '四、共性、差异及合肥工业大学现行方案情况')
add_body(doc, '从共同设置看，国内7校均保留财务会计、管理会计、审计、财务管理和会计信息系统等专业内容，同时设置计算机、信息系统或数据分析基础课程。国外3个项目均包含跨学科或宽口径课程，南洋理工大学和伊利诺伊大学还设置了计入培养要求的结构化实践。国内外样本均在会计专业课程之外安排数据、数量方法或信息技术内容。')
add_body(doc, '样本之间的差异主要体现在课程组织和实践形式。国内高校通常以总学分、课程模块和集中实践组织培养方案，主课程体系比较值为140—163学分；国外3校分别采用AU、LSE unit和semester hour，其学分口径与国内不同。南洋理工大学设置30周Work-Study，伊利诺伊大学设置3学分真实客户课程，伦敦政治经济学院官网项目页未列出必修长周期实习。国内样本在专业方向数量、课外学分和实践周数等方面也采用了不同安排。')
add_body(doc, '合肥工业大学2023版方案共165学分，其中理论课程134.5学分、实践教学30.5学分，分别占81.5%和18.5%。专业选修课程池共54学分，最低修读34.5学分，占总学分的20.9%；另设置2学分课外科研训练，不计入165学分。按本报告统计口径，现行方案总学分比7校平均值153.9学分多11.1学分，比样本最高值163学分多2学分。课程体系中已设置C++、数据库、人工智能、Python、大数据分析与商业智能、数据建模与智能财务决策、ERP和RPA等课程。')

add_h1(doc, '五、调研结果与培养方案修订的衔接')
add_body(doc, '结合国内外样本和合肥工业大学现行方案，本部分按照培养目标、毕业要求、课程体系、实践教学、人才培养模式和教育教学改革六个方面，对调研结果与后续修订工作进行对应梳理。')
add_subpara(doc, '（一）培养目标。', '国内外样本多将会计专业能力、数据分析能力和业务理解纳入培养目标。结合合肥工业大学的办学背景，培养目标修订可围绕先进制造业和数字经济场景，进一步梳理企业核算、分析、控制和决策等岗位任务。制造业企业和毕业校友访谈可用于补充业务财务、智能财务分析、审计内控等岗位信息，Python、RPA等具体工具名称可在课程说明中体现。')
add_subpara(doc, '（二）毕业要求。', '样本中的毕业要求主要包括伦理责任、专业胜任、业务理解、数据技术、问题解决与决策、研究创新、沟通协作和持续学习等内容。课程与毕业要求之间的对应关系，可结合课程项目、实习、团队展示、企业导师评价和毕业成果进行梳理。达成度评价中的个人达标线和班级达标比例，可在试运行基础上结合实际数据确定。')
add_subpara(doc, '（三）课程体系。', '合肥工业大学现行方案已设置C++、Python、数据库、人工智能、大数据分析与商业智能、数据建模与智能财务决策、ERP和RPA等课程。后续课程盘点可重点梳理管理统计学、计量经济学、管理决策分析、大数据分析与商业智能，以及智能财务管理、智能财务分析、数据建模与智能财务决策之间的内容衔接。结合校级公共课程是否同步调整，可分别测算160学分和156学分两种课程结构。编程和数据库内容可结合财务分析、审计和信息系统课程中的实际任务进行安排。')
add_subpara(doc, '（四）实践教学。', '国内外样本中的实践教学主要包括课程实验、案例项目、企业实习、真实客户任务和毕业成果等形式。合肥工业大学现有实践教学为30.5学分，后续可在现有总量内梳理独立实践和课内实践的对应关系。实践内容可按学年展开：第一学年侧重企业流程认知、岗位访谈和基础数据工具；第二学年侧重会计项目及ERP、RPA流程；第三学年侧重智能财务决策、审计与内控案例；第四学年衔接企业实习和毕业成果。校外项目的学分认定可结合任务书、过程记录、学生成果和企业导师评价。')
add_subpara(doc, '（五）人才培养模式。', '样本中较常见的组织方式是在共同会计核心课程之外设置专业方向或跨学科模块。结合现行34.5学分专业选修要求，可进一步梳理数智会计与智能决策、公司财务与资本市场、审计风险与可持续发展等方向课程，并与跨学院课程、微专业和研究训练衔接。方向课程的具体规模可结合选课情况、师资安排和岗位调研结果确定。')
add_subpara(doc, '（六）教育教学改革举措。', '国内外样本中的改革内容主要涉及课程更新、项目教学、校企协同和学习成果评价。结合培养方案修订，可同步梳理新设课程与现有课程的关系、先修要求和成果形式，并按学年汇总课程作品、项目、实习和毕业成果。相关材料可作为课程内容调整和培养方案后续修订的依据。')

add_caption(doc, '表3  培养方案修订工作安排')
table3_data = [
['工作','时间','拟形成材料','负责人'],
['课程盘点','0—3个月','课程知识点、先修关系、实践任务和重复内容清单','各课程组'],
['结构设计','3—6个月','160/156学分两套测算、共同核心、3个方向包和项目链','专业负责人、系务委员会'],
['课程开发','6—12个月','制造业案例、脱敏数据集、项目任务书和统一评分量表','课程团队、企业导师'],
['试运行与反馈','第1学年及以后','选取一个年级试点；每学年形成达成度报告和课程调整清单','学院教学委员会']]
t3 = doc.add_table(rows=len(table3_data), cols=4)
t3.alignment = WD_TABLE_ALIGNMENT.CENTER; t3.autofit=False
w3=[Cm(2.5),Cm(2.6),Cm(8.2),Cm(3.6)]
for i,row in enumerate(t3.rows):
    for j,cell in enumerate(row.cells):
        cell.width=w3[j]; cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=2 else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(0)
        rr=p.add_run(table3_data[i][j]); set_east_asia_font(rr,'宋体',8.5,i==0)
        set_cell_border(cell, top={'val':'single','sz':'6','color':'000000'}, bottom={'val':'single','sz':'6','color':'000000'}, left={'val':'single','sz':'4','color':'000000'}, right={'val':'single','sz':'4','color':'000000'})
set_repeat_table_header(t3.rows[0])
add_note(doc, '注：表中时间从修订工作正式启动之日起计算。试运行材料可在首个学年结束后统一汇总。')

add_h1(doc, '六、结语')
add_body(doc, '本报告对国内外样本的培养目标、毕业要求、课程体系、实践教学、人才培养模式和教育教学改革举措进行了整理，并呈现了样本之间的共同设置、不同做法以及合肥工业大学现行方案的基本情况。相关调研结果可与学校培养方案修订要求、课程组意见以及企业和校友反馈一并使用。')

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
r=p.add_run('主要资料来源'); set_east_asia_font(r,'黑体',11,False)
refs = [
'[1] 合肥工业大学管理学院：《合肥工业大学2023版会计学专业人才培养方案》，https://som.hfut.edu.cn/info/1035/7573.htm，访问日期：2026-08-02。',
'[2] 国内样本培养方案：北京理工大学（2019）、北京航空航天大学（2022）、华中科技大学（2021）、厦门大学四方向（2019）、哈尔滨工业大学（2022）、浙江大学（2019）和中南大学（2023）；各校主课程体系比较值及另列要求见表1。',
'[3] 中南大学商学院：《会计与财务系为会计21级本科生举行培养方案宣讲》，https://bs.csu.edu.cn/info/1046/14253.htm，访问日期：2026-08-02。',
'[4] London School of Economics and Political Science, BSc Accounting and Finance, https://www.lse.ac.uk/study-at-lse/undergraduate/bsc-accounting-and-finance，访问日期：2026-08-02。',
'[5] University of Illinois Urbana-Champaign, 2026–2027 Course Catalog: Accountancy + Data Science, BS, https://catalog.illinois.edu/undergraduate/bus/accountancy-data-science-bs/，访问日期：2026-08-02。',
'[6] Nanyang Technological University, Accountancy for Future Leaders—Bachelor of Accountancy in Sustainability Management and Analytics, https://www.ntu.edu.sg/engineering/coe-programmes/undergraduate/coe-programme-detail/accountancy-for-future-leaders-bachelor-of-accountancy-in-sustainability-management-and-analytics，访问日期：2026-08-02。']
for ref in refs:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(ref); set_east_asia_font(r,'宋体',8.5)

doc.save(OUT)
print(OUT)
