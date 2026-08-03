from pathlib import Path
import os
import subprocess
import sys

subprocess.run([sys.executable, '-m', 'pip', 'install', '-q', 'python-docx', 'matplotlib'], check=True)

import matplotlib.pyplot as plt
from matplotlib import font_manager
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('tmp_word_report/合肥工业大学会计学专业培养方案对标调研分析报告_语言修订版.docx')
CHART = Path('tmp_word_report/figure1.png')
OUT.parent.mkdir(parents=True, exist_ok=True)

# 单色图表，保持原报告中的学分位置图。
labels = ['HFUT', 'BUAA', 'CSU', 'HUST', 'ZJU', 'BIT', 'HIT', 'XMU']
values = [165, 163, 159.5, 157.5, 155, 152, 150, 140]
plt.figure(figsize=(8.2, 4.3))
plt.barh(labels[::-1], values[::-1], color='black')
plt.xlim(135, 168)
plt.xlabel('Credits')
for y, v in enumerate(values[::-1]):
    plt.text(v + 0.35, y, f'{v:g}', va='center', fontsize=9)
plt.grid(axis='x', linestyle=':', linewidth=0.6, alpha=0.65)
plt.tight_layout()
plt.savefig(CHART, dpi=220, bbox_inches='tight', facecolor='white')
plt.close()


def set_run_font(run, east='仿宋_GB2312', size=12, bold=False, latin='Times New Roman'):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_cell_margins(cell, top=70, start=70, bottom=70, end=70):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, val in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, size='6'):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = borders.find(qn(f'w:{edge}'))
        if tag is None:
            tag = OxmlElement(f'w:{edge}')
            borders.append(tag)
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), size)
        tag.set(qn('w:space'), '0')
        tag.set(qn('w:color'), '000000')


def format_body_paragraph(p, indent=True, after=0):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    if indent:
        pf.first_line_indent = Cm(0.85)


def add_body(text, indent=True):
    p = doc.add_paragraph()
    format_body_paragraph(p, indent=indent)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_h1(text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, east='黑体', size=12, bold=False)
    return p


def add_inline_subsection(label, text):
    p = doc.add_paragraph()
    format_body_paragraph(p, indent=False)
    r1 = p.add_run(label)
    set_run_font(r1, east='黑体', size=12, bold=False)
    r2 = p.add_run(text)
    set_run_font(r2)
    return p


def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, east='宋体', size=10.5)
    return p


def format_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths and ci < len(widths):
                cell.width = Cm(widths[ci])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 or ci < 3 else WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r, east='宋体', size=9, bold=(ri == 0))


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.3)
sec.header_distance = Cm(1.0)
sec.footer_distance = Cm(1.0)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), '仿宋_GB2312')
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)

# 标题
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(12)
p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
r = p.add_run('合肥工业大学会计学专业培养方案对标调研分析报告')
set_run_font(r, east='方正小标宋简体', size=18, bold=False)

add_h1('一、调研范围、比较口径与数据处理')
add_body('本次调研围绕培养目标、毕业要求、课程体系、实践教学、人才培养模式和教育教学改革举措六个方面展开。国内选取北京理工大学、北京航空航天大学、华中科技大学、厦门大学、哈尔滨工业大学、浙江大学和中南大学，共涉及7所高校、10份会计学相关培养方案，方案版本集中在2019—2023年。其中，厦门大学包括会计学、注册会计师、国际会计和CIMA四个方向，因此学校数量按7所统计，文档数量按10份统计。国外分别选取亚洲、欧洲和美洲各1所高校，即新加坡南洋理工大学、英国伦敦政治经济学院和美国伊利诺伊大学厄巴纳—香槟分校。相关资料主要来自学校官网项目页面和2026—2027课程目录，检索截至2026年8月2日。')
add_body('国内高校的学分统计采用主课程体系比较值。培养方案正文或课程结构主表明确列入的课内课程和实践教学学分计入比较值；另行列出的第二课堂、课外活动和附加模块，如未计入主表合计，则单独说明；已经包含在主表合计中的项目不再重复计算。表1列出了各校的取值依据。统计以学校为单位，厦门大学四个方向在7校统计中计1次。国外高校采用的LSE unit、semester hour和AU分别属于不同学分制度，本报告仅呈现各项目内部的课程结构，不作直接换算。')

add_caption('表1  国内7校会计学相关方案主课程体系学分比较值及认定依据')
t1_data = [
    ['学校', '版本', '比较值', '比较值认定依据（学分）'],
    ['北京航空航天大学', '2022', '163', '数字化与智能会计方向正文规定毕业总学分163；比较值取163'],
    ['中南大学', '2023', '159.5', '课程结构主表合计159.5，已含课外研学4；比较值取159.5'],
    ['华中科技大学', '2021', '157.5', '课内课程与集中实践合计157.5；另要求课外5，后者不计入'],
    ['浙江大学', '2019', '155', '文件写作155+5.5+6+8；155为主课程体系，另列19.5不计入比较值'],
    ['北京理工大学', '2019', '152', '正文规定最低152，课程计划表合计149；以正文为准，比较值取152'],
    ['哈尔滨工业大学', '2022', '150', '正文规定修满150并完成论文答辩方可毕业；比较值取150'],
    ['厦门大学', '2019', '140', '会计学、注册会计师、国际会计、CIMA四方向均为140；学校统计只计1次，比较值取140'],
]
t1 = doc.add_table(rows=len(t1_data), cols=4)
for i, row in enumerate(t1_data):
    for j, val in enumerate(row):
        t1.cell(i, j).text = val
format_table(t1, widths=[3.4, 1.6, 1.8, 9.2])

add_body('注：比较值用于本次样本均值、中位数和四分位数的计算，不涵盖学生全部毕业条件。浙江大学另列的19.5学分由“+5.5学分”、跨专业与国际化模块6学分、第二至第四课堂8学分构成；华中科技大学另列课外5学分。相关数据均根据各校培养方案正文和课程结构表整理。', indent=False)
add_body('按照上述口径，7校主课程体系比较值合计1077学分，平均值为153.9学分，中位数为155学分。比较值最低为140学分，最高为163学分，极差为23学分。采用线性插值计算，第一四分位数为151.0学分，第三四分位数为158.5学分，中间50%的学校分布在151.0—158.5学分之间。按本次7校样本总体计算，标准差为7.0学分。如按10份文档统计，厦门大学四个方向会形成重复权重，平均值为149.7学分，中位数为151学分。因此，后文统一采用7校统计口径。')

add_caption('图1  国内7校及合肥工业大学主课程体系学分位置')
picp = doc.add_paragraph()
picp.alignment = WD_ALIGN_PARAGRAPH.CENTER
picp.add_run().add_picture(str(CHART), width=Cm(14.8))
add_body('注：合肥工业大学现行方案共165学分，其中理论课程134.5学分、实践教学30.5学分，与7校主课程体系比较值采用同类统计范围。按该口径计算，合肥工业大学比7校平均值153.9学分多11.1学分，相差7.2%；比样本最高值163学分多2学分。', indent=False)

add_h1('二、国内高校六维调研结果')
add_inline_subsection('（一）培养目标。', '国内样本的培养目标主要涉及会计专业知识、数据分析和决策支持。北京航空航天大学将信息技术、大数据分析和智能技术纳入培养目标，提出运用财务、业务和社会经济数据支持管理决策；华中科技大学大数据与商业分析方向同时列出会计审计知识、程序设计和数据分析能力。两所学校均结合业务问题和决策任务描述相关能力。')
add_inline_subsection('（二）毕业要求。', '国内样本中的毕业要求通常涵盖专业知识、问题分析、实践应用、沟通协作、职业规范、创新和持续学习等内容。华中科技大学列出程序设计、管理信息系统开发和大数据分析方法；北京航空航天大学将运用计算机和大数据处理技术解决会计实务与决策问题纳入毕业要求。部分方案还将课程项目、实习和毕业成果作为相关能力的评价载体。')
add_inline_subsection('（三）课程体系。', '财务会计、管理会计、成本会计、审计、财务管理、税法和会计信息系统是国内样本中较为稳定的专业核心课程。以7所学校为统计单位，7所学校均开设计算机、信息系统或数据分析基础课程，占100%；北京理工大学、北京航空航天大学、华中科技大学、浙江大学和中南大学明确开设Python、C或C++课程，共5所，占71.4%。在程序设计课程之后，部分学校继续设置数据库、业务分析或智能会计课程。')
add_inline_subsection('（四）实践教学。', '不同学校采用不同方式组织和计量实践教学。北京航空航天大学的社会课堂（生产实习）为5学分、320学时；华中科技大学集中性实践教学为17学分、34周；哈尔滨工业大学的生产实习和毕业实习各3学分、各3周。相关方案分别以学分、学时和周数说明实践要求，对实践任务、企业参与和成果形式的列示方式也有所不同。')
add_inline_subsection('（五）人才培养模式。', '厦门大学在共同基础课程之外设置会计学、注册会计师、国际会计和CIMA四个方向，四个方向均为140学分；北京航空航天大学设置数字化与智能会计方向；浙江大学的课程结构包括宽口径基础、专业选修、跨专业模块和国际化模块。各校在共同核心课程、专业方向和跨学科选修空间方面采用了不同的组织方式。')
add_inline_subsection('（六）教育教学改革举措。', '从国内样本看，相关改革主要涉及数智课程建设、案例与项目教学、校企合作和课程评价调整。部分学校按照程序设计、数据库或信息系统、大数据分析、智能财务或智能决策的顺序安排课程；部分学校通过专题课程或方向选修课程呈现相关内容。不同方案对课程先修关系和综合项目的安排各有侧重。')

add_h1('三、国外三所高校调研情况')
add_body('国外样本分别选取亚洲、欧洲和美洲各1所高校。南洋理工大学的项目资料主要呈现职业资格与长周期实习的衔接，伦敦政治经济学院的课程结构包括社会科学和数量方法，伊利诺伊大学厄巴纳—香槟分校将会计课程与数据科学课程组合设置。表2沿用各校官方学分口径，相关数字仅用于说明项目内部结构。')
add_caption('表2  亚洲、欧洲、美洲三所高校官方项目数据与培养机制')
t2_data = [
    ['区域', '院校与项目', '官方数据', '培养机制'],
    ['亚洲', '南洋理工大学\nAccountancy for Future Leaders', '学制4年，授予荣誉学士；总计140 AU；30周Work-Study计15 AU', '会计核心与可持续管理、分析能力一体化；实习机构须获ACRA认可，并衔接SCAQ部分考试减免'],
    ['欧洲', '伦敦政治经济学院\nBSc Accounting and Finance', '3年修读12 LSE units，另修LSE100；项目获ACCA、CIMA、ICAEW认可用于部分考试减免', '会计与金融约占一半，配套经济学、数学、统计、计量经济学及校外选修；官方项目页未列必修长周期实习'],
    ['美洲', '伊利诺伊大学厄巴纳—香槟分校\nAccountancy + Data Science', '毕业最低124 semester hours；商科核心42、会计要求21、数据科学核心29—30；真实客户综合课程3学分', '会计系与统计、计算机、信息学院和数学系协同；数据科学伦理、算法、建模和数据治理贯穿课程'],
]
t2 = doc.add_table(rows=len(t2_data), cols=4)
for i, row in enumerate(t2_data):
    for j, val in enumerate(row):
        t2.cell(i, j).text = val
format_table(t2, widths=[1.2, 4.0, 5.0, 6.0])
add_body('注：AU、LSE unit和semester hour分别是南洋理工大学、伦敦政治经济学院和美国高校使用的学分单位，三者不作直接换算。ACRA指新加坡会计与企业监管局，SCAQ指新加坡特许会计师资格考试。伊利诺伊大学124学时还包括通识教育等要求，表列模块不与中国高校学分直接比较。', indent=False)
add_inline_subsection('（一）亚洲：南洋理工大学。', '该项目将可持续管理、分析课程、专业会计和职业实践纳入同一培养路径。30周Work-Study实习计15 AU，实习单位须经新加坡会计与企业监管局认可，相关经历可计入新加坡特许会计师资格要求的实践年限。项目资料同时列明实习机构条件、学分认定和职业资格衔接方式。')
add_inline_subsection('（二）欧洲：伦敦政治经济学院。', '项目学制为3年，学生修读12个LSE units，并完成LSE100。课程除会计和金融外，还包括经济学、数学、统计和计量经济学，学生可按规定选修校外课程。ACCA、CIMA和ICAEW的部分考试减免与学生的具体选课组合相衔接。')
add_inline_subsection('（三）美洲：伊利诺伊大学厄巴纳—香槟分校。', '该项目毕业最低要求为124学时，其中会计课程21学时、数据科学核心课程29—30学时。数据科学部分包括数学基础、算法与数据结构、数据伦理和数据治理。BUS 301为3学分真实客户综合课程，课程内容包括商业问题处理、数据分析和客户展示。项目分别规定了会计课程、数据科学课程和综合课程的学分要求。')
add_body('按本报告选取的3所学校统计，3校均采用跨学科或宽口径课程结构，占100%；南洋理工大学和伊利诺伊大学将数据分析或数据科学纳入课程体系，占66.7%；南洋理工大学和伦敦政治经济学院明确提供职业资格考试减免，占66.7%；南洋理工大学的30周实习和伊利诺伊大学的3学分真实客户课程均计入培养要求，结构化实践占66.7%。其中，南洋理工大学设置长周期必修实习，占33.3%。上述比例仅反映本次3个案例的基本情况。')

add_h1('四、共性、差异及合肥工业大学现行方案情况')
add_body('从共同设置看，国内7校均保留财务会计、管理会计、审计、财务管理和会计信息系统等专业内容，同时设置计算机、信息系统或数据分析基础课程。国外3个项目均包含跨学科或宽口径课程，南洋理工大学和伊利诺伊大学还设置了计入培养要求的结构化实践。国内外样本均在会计专业课程之外安排数据、数量方法或信息技术内容。')
add_body('样本之间的差异主要体现在课程组织和实践形式。国内高校通常以总学分、课程模块和集中实践组织培养方案，主课程体系比较值为140—163学分；国外3校分别采用AU、LSE unit和semester hour，其学分口径与国内不同。南洋理工大学设置30周Work-Study，伊利诺伊大学设置3学分真实客户课程，伦敦政治经济学院官网项目页未列出必修长周期实习。国内样本在专业方向数量、课外学分和实践周数等方面也采用了不同安排。')
add_body('合肥工业大学2023版方案共165学分，其中理论课程134.5学分、实践教学30.5学分，分别占81.5%和18.5%。专业选修课程池共54学分，最低修读34.5学分，占总学分的20.9%；另设置2学分课外科研训练，不计入165学分。按本报告统计口径，现行方案总学分比7校平均值153.9学分多11.1学分，比样本最高值163学分多2学分。课程体系中已设置C++、数据库、人工智能、Python、大数据分析与商业智能、数据建模与智能财务决策、ERP和RPA等课程。')

add_h1('五、调研结果对培养方案修订的启示')
add_body('结合国内外样本和合肥工业大学现行方案，调研结果主要涉及培养目标、毕业要求、课程体系、实践教学、人才培养模式和教育教学改革六个方面。相关内容用于说明不同高校的具体做法，并为后续修订工作提供资料参考。')
add_inline_subsection('（一）培养目标。', '国内外样本多将会计专业能力、数据分析能力和业务理解纳入培养目标。结合合肥工业大学的办学背景，培养目标修订可围绕先进制造业和数字经济场景，进一步梳理企业核算、分析、控制和决策等岗位任务。制造业企业和毕业校友访谈可补充业务财务、智能财务分析、审计内控等岗位信息，Python、RPA等具体工具名称可在课程说明中体现。')
add_inline_subsection('（二）毕业要求。', '样本中的毕业要求主要包括伦理责任、专业胜任、业务理解、数据技术、问题解决与决策、研究创新、沟通协作和持续学习等内容。课程与毕业要求之间的对应关系，可结合课程项目、实习、团队展示、企业导师评价和毕业成果进行梳理。达成度评价中的个人达标线和班级达标比例，可在试运行基础上结合实际数据确定。')
add_inline_subsection('（三）课程体系。', '合肥工业大学现行方案已设置C++、Python、数据库、人工智能、大数据分析与商业智能、数据建模与智能财务决策、ERP和RPA等课程。后续课程盘点可重点梳理管理统计学、计量经济学、管理决策分析、大数据分析与商业智能，以及智能财务管理、智能财务分析、数据建模与智能财务决策之间的内容衔接。结合校级公共课程是否同步调整，可分别测算160学分和156学分两种课程结构。编程和数据库内容可结合财务分析、审计和信息系统课程中的实际任务进行安排。')
add_inline_subsection('（四）实践教学。', '国内外样本中的实践教学主要包括课程实验、案例项目、企业实习、真实客户任务和毕业成果等形式。合肥工业大学现有实践教学为30.5学分，后续可在现有总量内梳理独立实践和课内实践的对应关系。实践内容可按学年展开：第一学年侧重企业流程认知、岗位访谈和基础数据工具；第二学年侧重会计项目及ERP、RPA流程；第三学年侧重智能财务决策、审计与内控案例；第四学年衔接企业实习和毕业成果。校外项目的学分认定可结合任务书、过程记录、学生成果和企业导师评价。')
add_inline_subsection('（五）人才培养模式。', '样本中较为常见的组织方式，是在共同会计核心课程之外设置专业方向或跨学科模块。结合现行34.5学分专业选修要求，可进一步梳理数智会计与智能决策、公司财务与资本市场、审计风险与可持续发展等方向课程，并与跨学院课程、微专业和研究训练相衔接。方向课程的具体规模，可结合学生选课情况、师资安排和岗位调研结果确定。')
add_inline_subsection('（六）教育教学改革举措。', '国内外样本中的改革内容主要涉及课程更新、项目教学、校企协同和学习成果评价。结合培养方案修订，可同步梳理新设课程与现有课程的关系、先修要求和成果形式，并按学年汇总课程作品、项目、实习和毕业成果。相关材料可用于记录课程运行情况，并作为后续课程调整和培养方案修订的参考。')

add_caption('表3  培养方案修订工作安排')
t3_data = [
    ['工作', '时间', '拟形成材料', '负责人'],
    ['课程盘点', '0—3个月', '课程知识点、先修关系、实践任务和重复内容清单', '各课程组'],
    ['结构设计', '3—6个月', '160/156学分两套测算、共同核心、3个方向包和项目链', '专业负责人、系务委员会'],
    ['课程开发', '6—12个月', '制造业案例、脱敏数据集、项目任务书和统一评分量表', '课程团队、企业导师'],
    ['试运行与反馈', '第1学年及以后', '选取一个年级试点；每学年形成达成度报告和课程调整清单', '学院教学委员会'],
]
t3 = doc.add_table(rows=len(t3_data), cols=4)
for i, row in enumerate(t3_data):
    for j, val in enumerate(row):
        t3.cell(i, j).text = val
format_table(t3, widths=[2.4, 2.5, 7.6, 3.5])
add_body('注：表中时间从修订工作正式启动之日起计算。试运行材料可在首个学年结束后统一汇总。', indent=False)

add_h1('六、结语')
add_body('本报告对国内外样本的培养目标、毕业要求、课程体系、实践教学、人才培养模式和教育教学改革举措进行了整理，并呈现了样本之间的共同设置、不同做法以及合肥工业大学现行方案的基本情况。相关调研结果可与学校培养方案修订要求、课程组意见以及企业和校友反馈一并使用。')

add_h1('主要资料来源')
sources = [
    '[1] 合肥工业大学管理学院：《合肥工业大学2023版会计学专业人才培养方案》，https://som.hfut.edu.cn/info/1035/7573.htm，访问日期：2026-08-02。',
    '[2] 国内样本培养方案：北京理工大学（2019）、北京航空航天大学（2022）、华中科技大学（2021）、厦门大学四方向（2019）、哈尔滨工业大学（2022）、浙江大学（2019）和中南大学（2023）；各校主课程体系比较值及另列要求见表1。',
    '[3] 中南大学商学院：《会计与财务系为会计21级本科生举行培养方案宣讲》，https://bs.csu.edu.cn/info/1046/14253.htm，访问日期：2026-08-02。',
    '[4] London School of Economics and Political Science, BSc Accounting and Finance, https://www.lse.ac.uk/study-at-lse/undergraduate/bsc-accounting-and-finance，访问日期：2026-08-02。',
    '[5] University of Illinois Urbana-Champaign, 2026–2027 Course Catalog: Accountancy + Data Science, BS, https://catalog.illinois.edu/undergraduate/bus/accountancy-data-science-bs/，访问日期：2026-08-02。',
    '[6] Nanyang Technological University, Accountancy for Future Leaders—Bachelor of Accountancy in Sustainability Management and Analytics, https://www.ntu.edu.sg/engineering/coe-programmes/undergraduate/coe-programme-detail/accountancy-for-future-leaders-bachelor-of-accountancy-in-sustainability-management-and-analytics，访问日期：2026-08-02。',
]
for s in sources:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    r = p.add_run(s)
    set_run_font(r, east='宋体', size=10.5)

# 页码
footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
run = fp.add_run()
run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)
set_run_font(run, east='宋体', size=9)

# 文档属性
props = doc.core_properties
props.title = '合肥工业大学会计学专业培养方案对标调研分析报告'
props.subject = '国内外高校会计学专业培养方案对标调研'
props.author = '合肥工业大学'

doc.save(OUT)
print(OUT)
